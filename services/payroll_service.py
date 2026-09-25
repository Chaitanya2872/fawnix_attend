"""Payroll service layer.

The payroll workbook supplied by the user is treated as a data template:
header columns become export columns and the salary breakup sheet becomes the
default calculation model.
"""

import calendar
import logging
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Dict, Iterable, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database.connection import get_db_connection, return_connection

logger = logging.getLogger(__name__)

GENERIC_ERROR = "Something went wrong handling this payroll request. Please try again."
MONEY_ZERO = Decimal("0.00")
MONEY_QUANT = Decimal("0.01")
DEFAULT_BONUS_RATE = Decimal("6.6628571429")
ESI_GROSS_LIMIT = Decimal("21000.00")

VALID_STATUSES = {"active", "inactive"}
VALID_ASSIGNMENT_STATUSES = {"draft", "active", "inactive"}
VALID_RUN_STATUSES = {"draft", "processed", "approved", "locked", "cancelled"}
VALID_COMPONENT_TYPES = {"earning", "employee_deduction", "employer_contribution", "informational"}
VALID_CALCULATION_TYPES = {"fixed", "percentage", "formula"}

HEADER_COLUMNS = [
    ("serial_no", "S#", "serial_no"),
    ("emp_code", "EMP ID", "emp_code"),
    ("project", "PROJECT", "project"),
    ("emp_name", "NAME", "emp_name"),
    ("designation", "DESIGNATION", "designation"),
    ("doj", "DOJ", "doj"),
    ("offer_type", "Offer Type", "offer_type"),
    ("offered_ctc", "offered CTC", "offered_ctc"),
    ("offered_gross_pay", "offered GP", "offered_gross_pay"),
    ("days_in_period", "Days", "days_in_period"),
    ("working_days", "No of working", "working_days"),
    ("basic_salary", "BS (50%)", "basic_salary"),
    ("hra", "HRA (50%)", "hra"),
    ("da", "DA", "da"),
    ("subtotal", "SUB TOTAL", "subtotal"),
    ("bonus", "Bonus", "bonus"),
    ("special_allowance", "Sp. Allow.", "special_allowance"),
    ("gross_pay", "GP", "gross_pay"),
    ("pf_employee", "PF 12%", "pf_employee"),
    ("esi_employee", "ESI 0.75%", "esi_employee"),
    ("professional_tax", "Proff Tax", "professional_tax"),
    ("salary_advance", "Sal Adv", "salary_advance"),
    ("gmi", "GMI", "gmi"),
    ("tds", "TDS", "tds"),
    ("total_deduction", "Total Deduction", "total_deduction"),
    ("net_pay", "NET PAY", "net_pay"),
    ("arrears", "Arrears", "arrears"),
    ("others", "Others", "others"),
    ("total_pay", "TOTAL PAY", "total_pay"),
    ("pf_employer", "PF 13%", "pf_employer"),
    ("esi_employer", "ESI@ 3.25", "esi_employer"),
    ("ctc", "CTC", "ctc"),
    ("remark", "Remark", "remark"),
]


def _money(value) -> Decimal:
    if value is None or value == "":
        return MONEY_ZERO
    if isinstance(value, Decimal):
        number = value
    else:
        try:
            number = Decimal(str(value).strip())
        except (InvalidOperation, ValueError) as exc:
            raise ValueError("amount must be a valid number") from exc
    return number.quantize(MONEY_QUANT, rounding=ROUND_HALF_UP)


def _percent(base: Decimal, rate: Decimal) -> Decimal:
    return _money((base * rate) / Decimal("100"))


def _serialize_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _serialize_row(row: Dict) -> Dict:
    return {key: _serialize_value(value) for key, value in (row or {}).items()}


def _normalize_text(value, *, required=False, field_name="field"):
    text = "" if value is None else str(value).strip()
    if required and not text:
        raise ValueError(f"{field_name} is required")
    return text or None


def _normalize_int(value, field_name: str):
    text = "" if value is None else str(value).strip()
    if not text:
        return None
    try:
        return int(text)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be a whole number") from exc


def _normalize_bool(value, default=False):
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"true", "1", "yes", "y"}:
        return True
    if normalized in {"false", "0", "no", "n"}:
        return False
    raise ValueError("boolean fields must be true or false")


def _normalize_date(value, field_name: str, *, required=False):
    text = "" if value is None else str(value).strip()
    if not text:
        if required:
            raise ValueError(f"{field_name} is required")
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise ValueError(f"{field_name} must be a valid date") from exc


def _pagination(params: Dict) -> Tuple[int, int, int]:
    page = max(int(params.get("page") or 1), 1)
    page_size = min(max(int(params.get("page_size") or 25), 1), 100)
    return page, page_size, (page - 1) * page_size


def _get_table_columns(cursor, table_name: str):
    cursor.execute(
        """
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = 'public'
          AND table_name = %s
        """,
        (table_name,),
    )
    return {row["column_name"] for row in (cursor.fetchall() or [])}


def _employee_select_expression(employee_columns, column_name: str, alias: str = None, cast: str = "TEXT"):
    select_alias = alias or column_name
    if column_name in employee_columns:
        return f"e.{column_name} AS {select_alias}"
    return f"NULL::{cast} AS {select_alias}"


def _employee_joined_date_select_expression(employee_columns):
    for column_name in ("emp_joined_date", "emp_joining_date"):
        if column_name in employee_columns:
            return f"e.{column_name} AS emp_joined_date"
    return "NULL::DATE AS emp_joined_date"


def _sort_clause(sort_by, sort_order, allowed, default):
    field = str(sort_by or default).strip()
    if field not in allowed:
        field = default
    order = str(sort_order or "asc").strip().lower()
    if order not in {"asc", "desc"}:
        order = "asc"
    return f"{field} {order.upper()}, id DESC"


def _normalize_component_payload(payload: Dict, *, partial=False):
    data = {}
    required = not partial
    if not partial or "component_code" in payload:
        data["component_code"] = _normalize_text(payload.get("component_code"), required=required, field_name="component_code")
    if not partial or "component_name" in payload:
        data["component_name"] = _normalize_text(payload.get("component_name"), required=required, field_name="component_name")
    if not partial or "component_type" in payload:
        component_type = _normalize_text(payload.get("component_type"), required=required, field_name="component_type")
        if component_type and component_type not in VALID_COMPONENT_TYPES:
            raise ValueError("component_type must be earning, employee_deduction, employer_contribution, or informational")
        data["component_type"] = component_type
    if not partial or "calculation_type" in payload:
        calculation_type = _normalize_text(payload.get("calculation_type")) or "fixed"
        if calculation_type not in VALID_CALCULATION_TYPES:
            raise ValueError("calculation_type must be fixed, percentage, or formula")
        data["calculation_type"] = calculation_type

    text_fields = ["base_component_code", "status", "created_by"]
    for field in text_fields:
        if field in payload:
            data[field] = _normalize_text(payload.get(field))
    if data.get("status") and data["status"] not in VALID_STATUSES:
        raise ValueError("status must be active or inactive")

    for field in ["rate_percent", "monthly_amount", "annual_amount", "cap_amount"]:
        if field in payload:
            data[field] = _money(payload.get(field)) if payload.get(field) not in (None, "") else None
    for field in ["is_taxable", "affects_gross", "affects_net", "affects_ctc"]:
        if field in payload:
            data[field] = _normalize_bool(payload.get(field))
    if "display_order" in payload:
        data["display_order"] = _normalize_int(payload.get("display_order"), "display_order") or 0

    return {key: value for key, value in data.items() if value is not None or key not in {"component_code", "component_name", "component_type"}}


def list_components(query_params: Dict):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        page, page_size, offset = _pagination(query_params)
        clauses = ["1=1"]
        values: List[object] = []
        search = str(query_params.get("search") or "").strip()
        if search:
            clauses.append("(component_code ILIKE %s OR component_name ILIKE %s)")
            values.extend([f"%{search}%", f"%{search}%"])
        component_type = str(query_params.get("component_type") or "").strip()
        if component_type:
            clauses.append("component_type = %s")
            values.append(component_type)
        status = str(query_params.get("status") or "").strip()
        if status:
            clauses.append("status = %s")
            values.append(status)
        where_clause = " AND ".join(clauses)
        cursor.execute(f"SELECT COUNT(*) AS total FROM payroll_components WHERE {where_clause}", values)
        total = int((cursor.fetchone() or {}).get("total") or 0)
        cursor.execute(
            f"""
            SELECT *
            FROM payroll_components
            WHERE {where_clause}
            ORDER BY {_sort_clause(query_params.get('sort_by'), query_params.get('sort_order'), {'component_code', 'component_name', 'component_type', 'display_order', 'status'}, 'display_order')}
            LIMIT %s OFFSET %s
            """,
            [*values, page_size, offset],
        )
        records = [_serialize_row(row) for row in (cursor.fetchall() or [])]
        return ({"success": True, "data": {"records": records, "count": total}}, 200)
    except Exception:
        logger.exception("Payroll component list failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def create_component(payload: Dict, created_by=None):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        record = _normalize_component_payload({**(payload or {}), "created_by": created_by or (payload or {}).get("created_by")})
        columns = list(record.keys())
        cursor.execute(
            f"""
            INSERT INTO payroll_components ({', '.join(columns)})
            VALUES ({', '.join(['%s'] * len(columns))})
            RETURNING *
            """,
            [record[column] for column in columns],
        )
        created = cursor.fetchone()
        conn.commit()
        return ({"success": True, "message": "Payroll component created successfully", "data": {"record": _serialize_row(created)}}, 201)
    except ValueError as exc:
        conn.rollback()
        return ({"success": False, "message": str(exc)}, 400)
    except Exception:
        conn.rollback()
        logger.exception("Payroll component create failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def update_component(component_id: int, payload: Dict):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        record = _normalize_component_payload(payload or {}, partial=True)
        record.pop("created_by", None)
        if not record:
            return ({"success": False, "message": "No valid fields provided"}, 400)
        assignments = [f"{column} = %s" for column in record]
        assignments.append("updated_at = NOW()")
        cursor.execute(
            f"""
            UPDATE payroll_components
            SET {', '.join(assignments)}
            WHERE id = %s
            RETURNING *
            """,
            [*[record[column] for column in record], component_id],
        )
        updated = cursor.fetchone()
        if not updated:
            conn.rollback()
            return ({"success": False, "message": "Payroll component not found"}, 404)
        conn.commit()
        return ({"success": True, "message": "Payroll component updated successfully", "data": {"record": _serialize_row(updated)}}, 200)
    except ValueError as exc:
        conn.rollback()
        return ({"success": False, "message": str(exc)}, 400)
    except Exception:
        conn.rollback()
        logger.exception("Payroll component update failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def delete_component(component_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM payroll_components WHERE id = %s RETURNING id", (component_id,))
        deleted = cursor.fetchone()
        if not deleted:
            conn.rollback()
            return ({"success": False, "message": "Payroll component not found"}, 404)
        conn.commit()
        return ({"success": True, "message": "Payroll component deleted successfully", "data": _serialize_row(deleted)}, 200)
    except Exception:
        conn.rollback()
        logger.exception("Payroll component delete failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def list_header_columns(query_params: Dict = None):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        status = str((query_params or {}).get("status") or "").strip()
        values: List[object] = []
        where_clause = "1=1"
        if status:
            where_clause = "status = %s"
            values.append(status)
        cursor.execute(
            f"""
            SELECT *
            FROM payroll_header_columns
            WHERE {where_clause}
            ORDER BY display_order ASC, id ASC
            """,
            values,
        )
        return ({"success": True, "data": {"records": [_serialize_row(row) for row in (cursor.fetchall() or [])]}}, 200)
    except Exception:
        logger.exception("Payroll header column list failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def salary_breakup_from_gross(gross_monthly, adjustments: Dict = None) -> Dict[str, Decimal]:
    adjustments = adjustments or {}
    gross = _money(gross_monthly)
    basic = _percent(gross, Decimal("50"))
    hra = _percent(basic, Decimal("50"))
    da = _percent(gross, Decimal("5"))
    bonus = _percent(gross, DEFAULT_BONUS_RATE)
    special_allowance = _money(gross - basic - hra - da - bonus)
    subtotal = _money(basic + hra + da)
    pf_employee = min(_percent(basic, Decimal("12")), Decimal("1800.00"))
    esi_employee = _percent(gross, Decimal("0.75")) if gross <= ESI_GROSS_LIMIT else MONEY_ZERO
    professional_tax = _money(adjustments.get("professional_tax", Decimal("200.00")))
    salary_advance = _money(adjustments.get("salary_advance"))
    gmi = _money(adjustments.get("gmi"))
    tds = _money(adjustments.get("tds"))
    total_deduction = _money(pf_employee + esi_employee + professional_tax + salary_advance + gmi + tds)
    net_pay = _money(gross - total_deduction)
    arrears = _money(adjustments.get("arrears"))
    others = _money(adjustments.get("others"))
    total_pay = _money(net_pay + arrears + others)
    pf_employer = min(_percent(basic, Decimal("13")), Decimal("1950.00"))
    esi_employer = _percent(gross, Decimal("3.25")) if gross <= ESI_GROSS_LIMIT else MONEY_ZERO
    ctc = _money(gross + pf_employer + esi_employer)

    return {
        "basic_salary": basic,
        "hra": hra,
        "da": da,
        "subtotal": subtotal,
        "bonus": bonus,
        "special_allowance": special_allowance,
        "gross_pay": gross,
        "pf_employee": pf_employee,
        "esi_employee": esi_employee,
        "professional_tax": professional_tax,
        "salary_advance": salary_advance,
        "gmi": gmi,
        "tds": tds,
        "total_deduction": total_deduction,
        "net_pay": net_pay,
        "arrears": arrears,
        "others": others,
        "total_pay": total_pay,
        "pf_employer": pf_employer,
        "esi_employer": esi_employer,
        "ctc": ctc,
    }


def salary_breakup_template(gross_monthly=None):
    gross = _money(gross_monthly or Decimal("35000.00"))
    breakup = salary_breakup_from_gross(gross)
    rows = [
        {"section": "earnings", "particulars": "Basic @ 50% of the Gross", "monthly": breakup["basic_salary"], "annual": breakup["basic_salary"] * 12},
        {"section": "earnings", "particulars": "HRA @ 50% of the Basic", "monthly": breakup["hra"], "annual": breakup["hra"] * 12},
        {"section": "earnings", "particulars": "DA", "monthly": breakup["da"], "annual": breakup["da"] * 12},
        {"section": "earnings", "particulars": "Bonus", "monthly": breakup["bonus"], "annual": breakup["bonus"] * 12},
        {"section": "earnings", "particulars": "Special Allowance", "monthly": breakup["special_allowance"], "annual": breakup["special_allowance"] * 12},
        {"section": "summary", "particulars": "Gross Salary (A)", "monthly": breakup["gross_pay"], "annual": breakup["gross_pay"] * 12},
        {"section": "employee_deductions", "particulars": "EPF", "monthly": breakup["pf_employee"], "annual": breakup["pf_employee"] * 12},
        {"section": "employee_deductions", "particulars": "ESIC", "monthly": breakup["esi_employee"], "annual": breakup["esi_employee"] * 12},
        {"section": "employee_deductions", "particulars": "Professional Tax", "monthly": breakup["professional_tax"], "annual": breakup["professional_tax"] * 12},
        {"section": "summary", "particulars": "Total Employee Ded. (B)", "monthly": breakup["total_deduction"], "annual": breakup["total_deduction"] * 12},
        {"section": "summary", "particulars": "Net Salary (Take Home) - C (A-B)", "monthly": breakup["net_pay"], "annual": breakup["net_pay"] * 12},
        {"section": "employer_deductions", "particulars": "EPF", "monthly": breakup["pf_employer"], "annual": breakup["pf_employer"] * 12},
        {"section": "employer_deductions", "particulars": "ESIC", "monthly": breakup["esi_employer"], "annual": breakup["esi_employer"] * 12},
        {"section": "summary", "particulars": "Total Employer Ded. (D)", "monthly": breakup["pf_employer"] + breakup["esi_employer"], "annual": (breakup["pf_employer"] + breakup["esi_employer"]) * 12},
        {"section": "summary", "particulars": "CTC (A+D)", "monthly": breakup["ctc"], "annual": breakup["ctc"] * 12},
    ]
    return {"success": True, "data": {"gross_monthly": str(gross), "rows": [{**row, "monthly": str(_money(row["monthly"])), "annual": str(_money(row["annual"]))} for row in rows]}}


def _normalize_structure_payload(payload: Dict, *, partial=False):
    data = {}
    required = not partial
    for field, required_flag in [
        ("structure_code", required),
        ("structure_name", required),
        ("offer_type", False),
        ("status", False),
        ("notes", False),
        ("created_by", False),
    ]:
        if not partial or field in payload:
            data[field] = _normalize_text(payload.get(field), required=required_flag, field_name=field)
    if data.get("status") and data["status"] not in VALID_STATUSES:
        raise ValueError("status must be active or inactive")
    if not partial:
        data.setdefault("status", "active")
    for field in ["payroll_unit_id"]:
        if field in payload:
            data[field] = _normalize_int(payload.get(field), field)
    for field in ["gross_monthly", "ctc_monthly"]:
        if field in payload:
            data[field] = _money(payload.get(field)) if payload.get(field) not in (None, "") else None
    for field in ["effective_from", "effective_to"]:
        if field in payload:
            data[field] = _normalize_date(payload.get(field), field)
    return {key: value for key, value in data.items() if value is not None or key in {"structure_code", "structure_name"}}


def _insert_structure_components(cursor, structure_id: int, components: Iterable[Dict]):
    for index, component in enumerate(components or [], start=1):
        code = _normalize_text(component.get("component_code"), required=True, field_name="component_code")
        name = _normalize_text(component.get("component_name"), required=True, field_name="component_name")
        component_type = _normalize_text(component.get("component_type"), required=True, field_name="component_type")
        if component_type not in VALID_COMPONENT_TYPES:
            raise ValueError("component_type must be earning, employee_deduction, employer_contribution, or informational")
        calculation_type = _normalize_text(component.get("calculation_type")) or "fixed"
        if calculation_type not in VALID_CALCULATION_TYPES:
            raise ValueError("calculation_type must be fixed, percentage, or formula")
        cursor.execute(
            """
            INSERT INTO payroll_salary_structure_components (
                structure_id, component_id, component_code, component_name, component_type,
                calculation_type, base_component_code, rate_percent, monthly_amount,
                annual_amount, cap_amount, display_order
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                structure_id,
                _normalize_int(component.get("component_id"), "component_id"),
                code,
                name,
                component_type,
                calculation_type,
                _normalize_text(component.get("base_component_code")),
                _money(component.get("rate_percent")) if component.get("rate_percent") not in (None, "") else None,
                _money(component.get("monthly_amount")) if component.get("monthly_amount") not in (None, "") else None,
                _money(component.get("annual_amount")) if component.get("annual_amount") not in (None, "") else None,
                _money(component.get("cap_amount")) if component.get("cap_amount") not in (None, "") else None,
                _normalize_int(component.get("display_order"), "display_order") or index,
            ),
        )


def list_salary_structures(query_params: Dict):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        page, page_size, offset = _pagination(query_params)
        clauses = ["1=1"]
        values: List[object] = []
        search = str(query_params.get("search") or "").strip()
        if search:
            clauses.append("(structure_code ILIKE %s OR structure_name ILIKE %s)")
            values.extend([f"%{search}%", f"%{search}%"])
        status = str(query_params.get("status") or "").strip()
        if status:
            clauses.append("status = %s")
            values.append(status)
        where_clause = " AND ".join(clauses)
        cursor.execute(f"SELECT COUNT(*) AS total FROM payroll_salary_structures WHERE {where_clause}", values)
        total = int((cursor.fetchone() or {}).get("total") or 0)
        cursor.execute(
            f"""
            SELECT *
            FROM payroll_salary_structures
            WHERE {where_clause}
            ORDER BY {_sort_clause(query_params.get('sort_by'), query_params.get('sort_order'), {'structure_code', 'structure_name', 'status', 'effective_from', 'created_at'}, 'structure_code')}
            LIMIT %s OFFSET %s
            """,
            [*values, page_size, offset],
        )
        return ({"success": True, "data": {"records": [_serialize_row(row) for row in (cursor.fetchall() or [])], "count": total}}, 200)
    except Exception:
        logger.exception("Salary structure list failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def _get_structure_components(cursor, structure_id: int):
    cursor.execute(
        """
        SELECT *
        FROM payroll_salary_structure_components
        WHERE structure_id = %s
        ORDER BY display_order ASC, id ASC
        """,
        (structure_id,),
    )
    return [_serialize_row(row) for row in (cursor.fetchall() or [])]


def get_salary_structure(structure_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM payroll_salary_structures WHERE id = %s", (structure_id,))
        record = cursor.fetchone()
        if not record:
            return ({"success": False, "message": "Salary structure not found"}, 404)
        data = _serialize_row(record)
        data["components"] = _get_structure_components(cursor, structure_id)
        return ({"success": True, "data": {"record": data}}, 200)
    except Exception:
        logger.exception("Salary structure read failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def create_salary_structure(payload: Dict, created_by=None):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        components = list((payload or {}).get("components") or [])
        record = _normalize_structure_payload({**(payload or {}), "created_by": created_by or (payload or {}).get("created_by")})
        columns = list(record.keys())
        cursor.execute(
            f"""
            INSERT INTO payroll_salary_structures ({', '.join(columns)})
            VALUES ({', '.join(['%s'] * len(columns))})
            RETURNING *
            """,
            [record[column] for column in columns],
        )
        created = cursor.fetchone()
        _insert_structure_components(cursor, int(created["id"]), components)
        conn.commit()
        data = _serialize_row(created)
        data["components"] = _get_structure_components(cursor, int(created["id"]))
        return ({"success": True, "message": "Salary structure created successfully", "data": {"record": data}}, 201)
    except ValueError as exc:
        conn.rollback()
        return ({"success": False, "message": str(exc)}, 400)
    except Exception:
        conn.rollback()
        logger.exception("Salary structure create failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def update_salary_structure(structure_id: int, payload: Dict):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        record = _normalize_structure_payload(payload or {}, partial=True)
        record.pop("created_by", None)
        components_supplied = "components" in (payload or {})
        if not record and not components_supplied:
            return ({"success": False, "message": "No valid fields provided"}, 400)
        if record:
            assignments = [f"{column} = %s" for column in record]
            assignments.append("updated_at = NOW()")
            cursor.execute(
                f"UPDATE payroll_salary_structures SET {', '.join(assignments)} WHERE id = %s RETURNING *",
                [*[record[column] for column in record], structure_id],
            )
            updated = cursor.fetchone()
            if not updated:
                conn.rollback()
                return ({"success": False, "message": "Salary structure not found"}, 404)
        else:
            cursor.execute("SELECT * FROM payroll_salary_structures WHERE id = %s", (structure_id,))
            updated = cursor.fetchone()
            if not updated:
                conn.rollback()
                return ({"success": False, "message": "Salary structure not found"}, 404)
        if components_supplied:
            cursor.execute("DELETE FROM payroll_salary_structure_components WHERE structure_id = %s", (structure_id,))
            _insert_structure_components(cursor, structure_id, payload.get("components") or [])
        conn.commit()
        data = _serialize_row(updated)
        data["components"] = _get_structure_components(cursor, structure_id)
        return ({"success": True, "message": "Salary structure updated successfully", "data": {"record": data}}, 200)
    except ValueError as exc:
        conn.rollback()
        return ({"success": False, "message": str(exc)}, 400)
    except Exception:
        conn.rollback()
        logger.exception("Salary structure update failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def delete_salary_structure(structure_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM payroll_salary_structures WHERE id = %s RETURNING id", (structure_id,))
        deleted = cursor.fetchone()
        if not deleted:
            conn.rollback()
            return ({"success": False, "message": "Salary structure not found"}, 404)
        conn.commit()
        return ({"success": True, "message": "Salary structure deleted successfully", "data": _serialize_row(deleted)}, 200)
    except Exception:
        conn.rollback()
        logger.exception("Salary structure delete failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def _normalize_assignment_payload(payload: Dict, *, partial=False):
    data = {}
    if not partial or "emp_code" in payload:
        data["emp_code"] = _normalize_text(payload.get("emp_code"), required=not partial, field_name="emp_code")
    for field in ["emp_email", "offer_type", "status", "notes", "created_by"]:
        if field in payload:
            data[field] = _normalize_text(payload.get(field))
    if data.get("status") and data["status"] not in VALID_ASSIGNMENT_STATUSES:
        raise ValueError("status must be draft, active, or inactive")
    if not partial:
        data.setdefault("status", "active")
    for field in ["structure_id", "payroll_unit_id"]:
        if field in payload:
            data[field] = _normalize_int(payload.get(field), field)
    for field in ["offered_ctc", "offered_gross_pay"]:
        if field in payload:
            data[field] = _money(payload.get(field)) if payload.get(field) not in (None, "") else None
    for field in ["effective_from", "effective_to"]:
        if field in payload:
            data[field] = _normalize_date(payload.get(field), field, required=(field == "effective_from" and not partial))
    if not partial:
        data.setdefault("effective_from", date.today())
    return {key: value for key, value in data.items() if value is not None or key == "emp_code"}


def list_employee_salary_assignments(query_params: Dict):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        page, page_size, offset = _pagination(query_params)
        clauses = ["1=1"]
        values: List[object] = []
        for field in ["emp_code", "status", "payroll_unit_id", "structure_id"]:
            raw = str(query_params.get(field) or "").strip()
            if raw:
                clauses.append(f"a.{field}::text = %s")
                values.append(raw)
        search = str(query_params.get("search") or "").strip()
        if search:
            clauses.append("(a.emp_code ILIKE %s OR a.emp_email ILIKE %s OR e.emp_full_name ILIKE %s)")
            values.extend([f"%{search}%", f"%{search}%", f"%{search}%"])
        where_clause = " AND ".join(clauses)
        cursor.execute(
            f"""
            SELECT COUNT(*) AS total
            FROM employee_salary_assignments a
            LEFT JOIN employees e ON e.emp_code = a.emp_code
            WHERE {where_clause}
            """,
            values,
        )
        total = int((cursor.fetchone() or {}).get("total") or 0)
        cursor.execute(
            f"""
            SELECT a.*, e.emp_full_name, e.emp_designation, e.emp_department, s.structure_code, s.structure_name
            FROM employee_salary_assignments a
            LEFT JOIN employees e ON e.emp_code = a.emp_code
            LEFT JOIN payroll_salary_structures s ON s.id = a.structure_id
            WHERE {where_clause}
            ORDER BY a.effective_from DESC, a.id DESC
            LIMIT %s OFFSET %s
            """,
            [*values, page_size, offset],
        )
        return ({"success": True, "data": {"records": [_serialize_row(row) for row in (cursor.fetchall() or [])], "count": total}}, 200)
    except Exception:
        logger.exception("Employee salary assignment list failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def create_employee_salary_assignment(payload: Dict, created_by=None):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        record = _normalize_assignment_payload({**(payload or {}), "created_by": created_by or (payload or {}).get("created_by")})
        if not record.get("structure_id") and not record.get("offered_gross_pay") and not record.get("offered_ctc"):
            raise ValueError("Provide structure_id, offered_gross_pay, or offered_ctc")
        columns = list(record.keys())
        cursor.execute(
            f"""
            INSERT INTO employee_salary_assignments ({', '.join(columns)})
            VALUES ({', '.join(['%s'] * len(columns))})
            RETURNING *
            """,
            [record[column] for column in columns],
        )
        created = cursor.fetchone()
        conn.commit()
        return ({"success": True, "message": "Employee salary assignment created successfully", "data": {"record": _serialize_row(created)}}, 201)
    except ValueError as exc:
        conn.rollback()
        return ({"success": False, "message": str(exc)}, 400)
    except Exception:
        conn.rollback()
        logger.exception("Employee salary assignment create failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def get_employee_salary_assignment(assignment_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            """
            SELECT a.*, e.emp_full_name, e.emp_designation, e.emp_department, s.structure_code, s.structure_name
            FROM employee_salary_assignments a
            LEFT JOIN employees e ON e.emp_code = a.emp_code
            LEFT JOIN payroll_salary_structures s ON s.id = a.structure_id
            WHERE a.id = %s
            """,
            (assignment_id,),
        )
        record = cursor.fetchone()
        if not record:
            return ({"success": False, "message": "Employee salary assignment not found"}, 404)
        return ({"success": True, "data": {"record": _serialize_row(record)}}, 200)
    except Exception:
        logger.exception("Employee salary assignment read failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def update_employee_salary_assignment(assignment_id: int, payload: Dict):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        record = _normalize_assignment_payload(payload or {}, partial=True)
        record.pop("created_by", None)
        if not record:
            return ({"success": False, "message": "No valid fields provided"}, 400)
        assignments = [f"{column} = %s" for column in record]
        assignments.append("updated_at = NOW()")
        cursor.execute(
            f"UPDATE employee_salary_assignments SET {', '.join(assignments)} WHERE id = %s RETURNING *",
            [*[record[column] for column in record], assignment_id],
        )
        updated = cursor.fetchone()
        if not updated:
            conn.rollback()
            return ({"success": False, "message": "Employee salary assignment not found"}, 404)
        conn.commit()
        return ({"success": True, "message": "Employee salary assignment updated successfully", "data": {"record": _serialize_row(updated)}}, 200)
    except ValueError as exc:
        conn.rollback()
        return ({"success": False, "message": str(exc)}, 400)
    except Exception:
        conn.rollback()
        logger.exception("Employee salary assignment update failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def delete_employee_salary_assignment(assignment_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM employee_salary_assignments WHERE id = %s RETURNING id", (assignment_id,))
        deleted = cursor.fetchone()
        if not deleted:
            conn.rollback()
            return ({"success": False, "message": "Employee salary assignment not found"}, 404)
        conn.commit()
        return ({"success": True, "message": "Employee salary assignment deleted successfully", "data": _serialize_row(deleted)}, 200)
    except Exception:
        conn.rollback()
        logger.exception("Employee salary assignment delete failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def _period_end(year: int, month: int):
    return date(year, month, calendar.monthrange(year, month)[1])


def _normalize_period(payload: Dict):
    year = _normalize_int(payload.get("period_year") or payload.get("year"), "period_year")
    month = _normalize_int(payload.get("period_month") or payload.get("month"), "period_month")
    if not year or not month:
        raise ValueError("period_year and period_month are required")
    if month < 1 or month > 12:
        raise ValueError("period_month must be between 1 and 12")
    return year, month


def _adjustment_for(payload: Dict, emp_code: str):
    adjustments = payload.get("adjustments") or {}
    if isinstance(adjustments, dict):
        return adjustments.get(emp_code) or adjustments.get(str(emp_code).lower()) or {}
    if isinstance(adjustments, list):
        for item in adjustments:
            if str(item.get("emp_code") or "").lower() == str(emp_code).lower():
                return item
    return {}


def _build_payroll_line(assignment: Dict, year: int, month: int, payload: Dict):
    adjustment = _adjustment_for(payload, assignment.get("emp_code"))
    days_in_period = _money(adjustment.get("days_in_period") or adjustment.get("days") or calendar.monthrange(year, month)[1])
    working_days = _money(adjustment.get("working_days") or adjustment.get("no_of_working") or days_in_period)
    gross = (
        assignment.get("offered_gross_pay")
        or assignment.get("structure_gross_monthly")
        or (Decimal(str(assignment.get("offered_ctc"))) / Decimal("12") if assignment.get("offered_ctc") else None)
    )
    if gross is None:
        raise ValueError(f"Gross pay is unavailable for {assignment.get('emp_code')}")
    breakup = salary_breakup_from_gross(gross, adjustment)
    return {
        "assignment_id": assignment.get("id"),
        "emp_code": assignment.get("emp_code"),
        "emp_email": assignment.get("emp_email") or assignment.get("employee_email"),
        "emp_name": assignment.get("emp_full_name") or assignment.get("emp_name"),
        "project": assignment.get("emp_department") or assignment.get("project"),
        "designation": assignment.get("emp_designation") or assignment.get("designation"),
        "doj": assignment.get("emp_joined_date") or assignment.get("emp_joining_date"),
        "offer_type": assignment.get("offer_type") or assignment.get("structure_offer_type"),
        "offered_ctc": _money(assignment.get("offered_ctc") or Decimal("0")),
        "offered_gross_pay": _money(gross),
        "days_in_period": days_in_period,
        "working_days": working_days,
        **breakup,
        "remark": _normalize_text(adjustment.get("remark") or adjustment.get("remarks")),
    }


def _component_rows_from_line(line: Dict):
    mapping = [
        ("basic_salary", "Basic @ 50% of the Gross", "earning", 10),
        ("hra", "HRA @ 50% of the Basic", "earning", 20),
        ("da", "DA", "earning", 30),
        ("bonus", "Bonus", "earning", 40),
        ("special_allowance", "Special Allowance", "earning", 50),
        ("pf_employee", "EPF", "employee_deduction", 110),
        ("esi_employee", "ESIC", "employee_deduction", 120),
        ("professional_tax", "Professional Tax", "employee_deduction", 130),
        ("pf_employer", "Employer EPF", "employer_contribution", 210),
        ("esi_employer", "Employer ESIC", "employer_contribution", 220),
    ]
    return [
        {
            "component_code": code,
            "component_name": name,
            "component_type": component_type,
            "monthly_amount": _money(line.get(code)),
            "annual_amount": _money(line.get(code)) * 12,
            "display_order": display_order,
        }
        for code, name, component_type, display_order in mapping
    ]


def _load_assignments(cursor, payload: Dict, year: int, month: int):
    payroll_unit_id = _normalize_int(payload.get("payroll_unit_id"), "payroll_unit_id") if payload.get("payroll_unit_id") else None
    employee_codes = payload.get("employee_codes") or []
    period_end = _period_end(year, month)
    employee_columns = _get_table_columns(cursor, "employees")
    employee_select_fields = [
        _employee_select_expression(employee_columns, "emp_full_name"),
        _employee_select_expression(employee_columns, "emp_email", "employee_email"),
        _employee_select_expression(employee_columns, "emp_department"),
        _employee_select_expression(employee_columns, "emp_designation"),
        _employee_joined_date_select_expression(employee_columns),
    ]
    clauses = [
        "a.status = 'active'",
        "a.effective_from <= %s",
        "(a.effective_to IS NULL OR a.effective_to >= %s)",
    ]
    values: List[object] = [period_end, date(year, month, 1)]
    if payroll_unit_id:
        clauses.append("a.payroll_unit_id = %s")
        values.append(payroll_unit_id)
    if employee_codes:
        placeholders = ", ".join(["%s"] * len(employee_codes))
        clauses.append(f"a.emp_code IN ({placeholders})")
        values.extend(employee_codes)
    employee_select_sql = ",\n            ".join(employee_select_fields)
    where_clause = " AND ".join(clauses)
    cursor.execute(
        f"""
        SELECT
            a.*,
            {employee_select_sql},
            s.gross_monthly AS structure_gross_monthly,
            s.ctc_monthly AS structure_ctc_monthly,
            s.offer_type AS structure_offer_type
        FROM employee_salary_assignments a
        LEFT JOIN employees e ON e.emp_code = a.emp_code
        LEFT JOIN payroll_salary_structures s ON s.id = a.structure_id
        WHERE {where_clause}
        ORDER BY a.emp_code ASC, a.effective_from DESC, a.id DESC
        """,
        values,
    )
    rows = cursor.fetchall() or []
    seen = set()
    unique_rows = []
    for row in rows:
        key = row.get("emp_code")
        if key in seen:
            continue
        seen.add(key)
        unique_rows.append(row)
    return unique_rows


def create_payroll_run(payload: Dict, created_by=None):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        year, month = _normalize_period(payload or {})
        assignments = _load_assignments(cursor, payload or {}, year, month)
        if not assignments:
            return ({"success": False, "message": "No active salary assignments found for this payroll run"}, 400)

        run_code = _normalize_text((payload or {}).get("run_code")) or f"PAY-{year}{month:02d}-{datetime.utcnow().strftime('%H%M%S')}"
        run_name = _normalize_text((payload or {}).get("run_name")) or f"Payroll {calendar.month_name[month]} {year}"
        payroll_unit_id = _normalize_int((payload or {}).get("payroll_unit_id"), "payroll_unit_id") if (payload or {}).get("payroll_unit_id") else None
        cursor.execute(
            """
            INSERT INTO payroll_runs (
                run_code, run_name, period_year, period_month, payroll_unit_id,
                status, processed_at, notes, created_by
            )
            VALUES (%s, %s, %s, %s, %s, 'processed', NOW(), %s, %s)
            RETURNING *
            """,
            (run_code, run_name, year, month, payroll_unit_id, _normalize_text((payload or {}).get("notes")), created_by),
        )
        run = cursor.fetchone()
        run_id = int(run["id"])
        totals = {"gross_total": MONEY_ZERO, "deduction_total": MONEY_ZERO, "net_total": MONEY_ZERO, "ctc_total": MONEY_ZERO}
        inserted_lines = []
        line_columns = [
            "run_id", "assignment_id", "emp_code", "emp_email", "emp_name", "project", "designation", "doj",
            "offer_type", "offered_ctc", "offered_gross_pay", "days_in_period", "working_days", "basic_salary",
            "hra", "da", "subtotal", "bonus", "special_allowance", "gross_pay", "pf_employee", "esi_employee",
            "professional_tax", "salary_advance", "gmi", "tds", "total_deduction", "net_pay", "arrears",
            "others", "total_pay", "pf_employer", "esi_employer", "ctc", "remark",
        ]
        for assignment in assignments:
            line = _build_payroll_line(assignment, year, month, payload or {})
            line["run_id"] = run_id
            cursor.execute(
                f"""
                INSERT INTO payroll_run_lines ({', '.join(line_columns)})
                VALUES ({', '.join(['%s'] * len(line_columns))})
                RETURNING *
                """,
                [line.get(column) for column in line_columns],
            )
            created_line = cursor.fetchone()
            inserted_lines.append(created_line)
            for component in _component_rows_from_line(line):
                cursor.execute(
                    """
                    INSERT INTO payroll_run_components (
                        run_line_id, component_code, component_name, component_type,
                        monthly_amount, annual_amount, display_order
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        created_line["id"],
                        component["component_code"],
                        component["component_name"],
                        component["component_type"],
                        component["monthly_amount"],
                        component["annual_amount"],
                        component["display_order"],
                    ),
                )
            totals["gross_total"] += _money(line["gross_pay"])
            totals["deduction_total"] += _money(line["total_deduction"])
            totals["net_total"] += _money(line["net_pay"])
            totals["ctc_total"] += _money(line["ctc"])

        cursor.execute(
            """
            UPDATE payroll_runs
            SET total_employees = %s,
                gross_total = %s,
                deduction_total = %s,
                net_total = %s,
                ctc_total = %s,
                updated_at = NOW()
            WHERE id = %s
            RETURNING *
            """,
            (
                len(inserted_lines),
                totals["gross_total"],
                totals["deduction_total"],
                totals["net_total"],
                totals["ctc_total"],
                run_id,
            ),
        )
        run = cursor.fetchone()
        conn.commit()
        return ({
            "success": True,
            "message": "Payroll run processed successfully",
            "data": {
                "run": _serialize_row(run),
                "lines": [_serialize_row(line) for line in inserted_lines],
            },
        }, 201)
    except ValueError as exc:
        conn.rollback()
        return ({"success": False, "message": str(exc)}, 400)
    except Exception:
        conn.rollback()
        logger.exception("Payroll run create failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def list_payroll_runs(query_params: Dict):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        page, page_size, offset = _pagination(query_params)
        clauses = ["1=1"]
        values: List[object] = []
        for field in ["period_year", "period_month", "status", "payroll_unit_id"]:
            raw = str(query_params.get(field) or "").strip()
            if raw:
                clauses.append(f"{field}::text = %s")
                values.append(raw)
        where_clause = " AND ".join(clauses)
        cursor.execute(f"SELECT COUNT(*) AS total FROM payroll_runs WHERE {where_clause}", values)
        total = int((cursor.fetchone() or {}).get("total") or 0)
        cursor.execute(
            f"""
            SELECT *
            FROM payroll_runs
            WHERE {where_clause}
            ORDER BY period_year DESC, period_month DESC, id DESC
            LIMIT %s OFFSET %s
            """,
            [*values, page_size, offset],
        )
        return ({"success": True, "data": {"records": [_serialize_row(row) for row in (cursor.fetchall() or [])], "count": total}}, 200)
    except Exception:
        logger.exception("Payroll run list failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def get_payroll_run(run_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM payroll_runs WHERE id = %s", (run_id,))
        run = cursor.fetchone()
        if not run:
            return ({"success": False, "message": "Payroll run not found"}, 404)
        cursor.execute("SELECT * FROM payroll_run_lines WHERE run_id = %s ORDER BY emp_code ASC", (run_id,))
        lines = [_serialize_row(row) for row in (cursor.fetchall() or [])]
        return ({"success": True, "data": {"run": _serialize_row(run), "lines": lines}}, 200)
    except Exception:
        logger.exception("Payroll run read failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def update_payroll_run_status(run_id: int, payload: Dict, actor=None):
    try:
        status = _normalize_text((payload or {}).get("status"), required=True, field_name="status")
        if status not in VALID_RUN_STATUSES:
            return ({"success": False, "message": "status must be draft, processed, approved, locked, or cancelled"}, 400)
    except ValueError as exc:
        return ({"success": False, "message": str(exc)}, 400)
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        approved_by = actor if status in {"approved", "locked"} else None
        cursor.execute(
            """
            UPDATE payroll_runs
            SET status = %s,
                approved_by = COALESCE(%s, approved_by),
                approved_at = CASE WHEN %s IS NOT NULL THEN NOW() ELSE approved_at END,
                notes = COALESCE(%s, notes),
                updated_at = NOW()
            WHERE id = %s
            RETURNING *
            """,
            (status, approved_by, approved_by, _normalize_text((payload or {}).get("notes")), run_id),
        )
        run = cursor.fetchone()
        if not run:
            conn.rollback()
            return ({"success": False, "message": "Payroll run not found"}, 404)
        conn.commit()
        return ({"success": True, "message": "Payroll run status updated successfully", "data": {"run": _serialize_row(run)}}, 200)
    except Exception:
        conn.rollback()
        logger.exception("Payroll run status update failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def get_payslip(run_id: int, emp_code: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM payroll_runs WHERE id = %s", (run_id,))
        run = cursor.fetchone()
        if not run:
            return ({"success": False, "message": "Payroll run not found"}, 404)
        cursor.execute(
            "SELECT * FROM payroll_run_lines WHERE run_id = %s AND LOWER(emp_code) = LOWER(%s)",
            (run_id, emp_code),
        )
        line = cursor.fetchone()
        if not line:
            return ({"success": False, "message": "Payslip not found"}, 404)
        cursor.execute(
            """
            SELECT *
            FROM payroll_run_components
            WHERE run_line_id = %s
            ORDER BY display_order ASC, id ASC
            """,
            (line["id"],),
        )
        components = [_serialize_row(row) for row in (cursor.fetchall() or [])]
        return ({"success": True, "data": {"run": _serialize_row(run), "line": _serialize_row(line), "components": components}}, 200)
    except Exception:
        logger.exception("Payslip read failed")
        return ({"success": False, "message": GENERIC_ERROR}, 500)
    finally:
        cursor.close()
        return_connection(conn)


def _rows_for_run_export(run: Dict, lines: List[Dict]):
    rows = []
    for index, line in enumerate(lines, start=1):
        row = {**line, "serial_no": index}
        rows.append(row)
    return rows


def build_payroll_run_workbook(run: Dict, lines: List[Dict]):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Payroll"
    headers = [label for _, label, _ in HEADER_COLUMNS]
    last_column = len(headers)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title = f"PAYROLL - {calendar.month_name[int(run.get('period_month') or 1)].upper()} {run.get('period_year')}"
    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="0D2B23")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28
    thin = Side(style="thin", color="B9C9CF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(2, col_index, header)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor="106B52")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_index, row in enumerate(_rows_for_run_export(run, lines), start=3):
        for col_index, (key, _label, source_field) in enumerate(HEADER_COLUMNS, start=1):
            value = row.get(source_field or key)
            cell = sheet.cell(row_index, col_index, _serialize_value(value))
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "L3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(last_column)}{max(2, len(lines) + 2)}"
    for index in range(1, last_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 13
    for index in [3, 4, 5, 33]:
        sheet.column_dimensions[get_column_letter(index)].width = 22
    sheet.sheet_view.showGridLines = False
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def export_payroll_run_workbook(run_id: int):
    response, status_code = get_payroll_run(run_id)
    if status_code != 200:
        return response, status_code
    run = response["data"]["run"]
    lines = response["data"]["lines"]
    return ({"success": True, "data": {"workbook": build_payroll_run_workbook(run, lines), "run": run}}, 200)


def build_payslip_workbook(run: Dict, line: Dict, components: List[Dict]):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Payslip"
    sheet["A1"] = "Payslip"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="0D2B23")
    sheet.merge_cells("A1:D1")
    meta = [
        ("Employee", line.get("emp_name")),
        ("Employee ID", line.get("emp_code")),
        ("Period", f"{calendar.month_name[int(run.get('period_month') or 1)]} {run.get('period_year')}"),
        ("Designation", line.get("designation")),
    ]
    for idx, (label, value) in enumerate(meta, start=3):
        sheet.cell(idx, 1, label).font = Font(bold=True)
        sheet.cell(idx, 2, value or "")
    start_row = 8
    sheet.cell(start_row, 1, "Particulars").font = Font(bold=True)
    sheet.cell(start_row, 2, "Type").font = Font(bold=True)
    sheet.cell(start_row, 3, "Monthly").font = Font(bold=True)
    sheet.cell(start_row, 4, "Annual").font = Font(bold=True)
    for index, component in enumerate(components, start=start_row + 1):
        sheet.cell(index, 1, component.get("component_name"))
        sheet.cell(index, 2, component.get("component_type"))
        sheet.cell(index, 3, component.get("monthly_amount"))
        sheet.cell(index, 4, component.get("annual_amount"))
    summary_row = start_row + len(components) + 2
    for offset, (label, value) in enumerate([
        ("Gross Pay", line.get("gross_pay")),
        ("Total Deduction", line.get("total_deduction")),
        ("Net Pay", line.get("net_pay")),
        ("CTC", line.get("ctc")),
    ]):
        sheet.cell(summary_row + offset, 1, label).font = Font(bold=True)
        sheet.cell(summary_row + offset, 2, value)
    for col in range(1, 5):
        sheet.column_dimensions[get_column_letter(col)].width = 22
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
